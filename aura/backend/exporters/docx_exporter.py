"""DOCX exporter."""
from docx import Document

from backend.exporters.base_exporter import BaseExporter


class DOCXExporter(
    BaseExporter
):

    def __init__(

        self

    ):

        super().__init__(

            ".docx"

        )

    def export(

        self,

        result,

        output_path

    ):

        self.validate(

            output_path

        )

        document = Document()

        document.add_heading(

            result.title,

            level=1

        )

        document.add_paragraph(

            result.content

        )

        document.save(

            output_path

        )