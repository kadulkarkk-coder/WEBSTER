"""Presentation exporter."""
from pptx import Presentation

from backend.exporters.base_exporter import BaseExporter


class PPTExporter(
    BaseExporter
):

    def __init__(

        self

    ):

        super().__init__(

            ".pptx"

        )

    def export(

        self,

        result,

        output_path

    ):

        self.validate(

            output_path

        )

        presentation = Presentation()

        layout = presentation.slide_layouts[1]

        slide = presentation.slides.add_slide(

            layout

        )

        slide.shapes.title.text = result.title

        slide.placeholders[1].text = result.content

        presentation.save(

            output_path

        )